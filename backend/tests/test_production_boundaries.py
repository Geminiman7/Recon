import os
os.environ.update(DATABASE_URL="sqlite://", SECRET_KEY="test-only-secret-with-at-least-32-bytes", ENVIRONMENT="test", REDIS_URL="", SENTRY_DSN="")

import asyncio
import hashlib
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID, uuid4
from unittest.mock import patch
from fastapi import Depends, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from openpyxl import Workbook, load_workbook
from starlette.requests import Request
from starlette.responses import Response
from starlette.datastructures import UploadFile

from app.main import app
from app.core.database import Base, get_db
from app.core.config import settings
from app.core.security import hash_password, create_access_token
from app.core.sessions import SESSION_COOKIE, CSRF_COOKIE, issue_csrf, validate_csrf, set_session
from app.api.dependencies import get_current_user, require_active_subscription
from app.models.company import Company
from app.models.user import User, UserRole
from app.models.job import ReconciliationJob, JobStatus
from app.models.reconciliation_result import ReconciliationResult, ReconciliationStatus
from app.models.password_reset_token import PasswordResetToken
from app.models.export_request import ExportRequest
from app.services.reconciliation_service import ReconciliationService
from app.services.auth_service import AuthService
from app.services.user_service import UserService
from app.services.export_service import generate_export, expire_exports
from app.services.export_writers import safe_cell, write_excel, write_pdf
from app.services.file_validation import validate_file
from app.services.upload_service import UploadService
from app.services import storage_service as storage


def active_user(user=Depends(get_current_user)):
    return user


class BoundaryTests(unittest.TestCase):
    def test_upload_mapping_reconciliation_results_workflow(self):
        from io import BytesIO
        from app.models.processor import Processor

        processor = Processor(name="Workflow processor", is_active=True)
        self.db.add(processor)
        self.db.commit()
        created = self.client.post("/jobs", json={"job_name": "Upload workflow"}, headers=self.csrf())
        self.assertEqual(created.status_code, 200, created.text)
        job_id = created.json()["id"]
        company = ("ID,Value,State\nmatched,100,SUCCESS\namount,200,SUCCESS\n"
                   "status,300,FAILED\ncompany-only,400,SUCCESS\n"
                   "duplicate,500,SUCCESS\nduplicate,500,SUCCESS\nlater-file,600,SUCCESS\n")
        first = ("Ref,Total,Outcome\nmatched,100,SUCCESS\namount,201,SUCCESS\n"
                 "status,300,SUCCESS\nprocessor-only,700,SUCCESS\nduplicate,500,SUCCESS\n")
        workbook = Workbook()
        workbook.active.append(["Ref", "Total", "Outcome"])
        workbook.active.append(["later-file", 600, "SUCCESS"])
        excel = BytesIO()
        workbook.save(excel)
        workbook.close()
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)), \
                patch.object(settings, "STORAGE_BACKEND", "local"), \
                patch.object(settings, "RECONCILIATION_MODE", "sync"):
            upload_ids = []
            for kind, filename, content, mime in (
                ("company", "company.csv", company.encode(), "text/csv"),
                ("processor", "first.csv", first.encode(), "text/csv"),
                ("processor", "second.xlsx", excel.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ):
                data = {"job_id": job_id}
                if kind == "processor":
                    data["processor_id"] = str(processor.id)
                response = self.client.post(f"/uploads/{kind}", data=data,
                    files={"file": (filename, content, mime)}, headers=self.csrf())
                self.assertEqual(response.status_code, 200, response.text)
                upload_id = response.json()["id"]
                upload_ids.append(upload_id)
                headers = self.client.get(f"/uploads/{upload_id}/columns")
                self.assertEqual(headers.status_code, 200, headers.text)
                self.assertEqual(len(headers.json()["columns"]), 3)
            listed = self.client.get(f"/uploads/job/{job_id}")
            self.assertEqual({row["id"] for row in listed.json()}, set(upload_ids))
            mapped = self.client.post("/mapping", headers=self.csrf(), json={
                "job_id": job_id, "company_file_id": upload_ids[0],
                "processor_file_ids": upload_ids[1:],
                "company": {"transaction_id": "ID", "amount": "Value", "status": "State"},
                "processors": {key: {"transaction_id": "Ref", "amount": "Total", "status": "Outcome"}
                               for key in upload_ids[1:]},
            })
            self.assertEqual(mapped.status_code, 200, mapped.text)
            self.assertEqual(mapped.json()["job_status"], "READY_TO_RECONCILE")
            for attempt in range(2):
                run = self.client.post(f"/reconciliation/{job_id}/run", headers=self.csrf())
                self.assertEqual(run.status_code, 200, run.text)
                self.assertEqual(run.json()["total"], 7)
                self.assertEqual(self.client.get(f"/jobs/{job_id}").json()["status"], "COMPLETED")
                results = self.client.get(f"/reconciliation/{job_id}/results?page_size=2")
                self.assertEqual(results.status_code, 200, results.text)
                summary = results.json()
                self.assertEqual((summary["total"], summary["matched"], summary["mismatched"], summary["missing"]), (7, 2, 3, 2))
                rows = []
                for page in range(1, summary["total_pages"] + 1):
                    rows.extend(self.client.get(f"/reconciliation/{job_id}/results?page_size=2&page={page}").json()["results"])
                self.assertEqual(len(rows), 7)
                self.assertEqual({row["status"] for row in rows}, {status.value for status in ReconciliationStatus})
                later = self.client.get(f"/reconciliation/{job_id}/results?search=later-file&status=MATCHED").json()
                self.assertEqual(later["filtered_total"], 1)
                self.assertEqual(later["results"][0]["difference"], 0)

    def test_missing_stored_files_return_actionable_errors_and_fail_attempt(self):
        from app.models.upload import Upload, UploadType
        from app.models.processor import Processor
        from app.models.column_mapping import ColumnMapping
        processor = Processor(name="Missing file test")
        self.db.add(processor)
        self.db.flush()
        records = []
        for kind, name in [(UploadType.COMPANY, "company.csv"), (UploadType.PROCESSOR, "processor.csv")]:
            record = Upload(company_id=self.company.id, job_id=self.job.id, uploaded_by=self.user.id,
                processor_id=processor.id if kind == UploadType.PROCESSOR else None,
                original_filename=name, stored_filename=name, storage_path="local://" + name,
                file_size=20, checksum=name, mime_type="text/csv", upload_type=kind)
            self.db.add(record)
            records.append(record)
        self.db.flush()
        self.db.add(ColumnMapping(company_id=self.company.id, job_id=self.job.id,
            processor_id=processor.id, processor_upload_id=records[1].id,
            company_column="transaction_id", processor_column="transaction_id", canonical_column="transaction_id"))
        self.db.commit()
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)), patch.object(settings, "RECONCILIATION_MODE", "sync"):
            for filename in ("company.csv", "processor.csv"):
                response = self.client.post(f"/reconciliation/{self.job.id}/run", headers=self.csrf())
                self.assertEqual(response.status_code, 422, response.text)
                self.assertIn(filename, response.json()["detail"])
                self.assertIn("new job", response.json()["detail"])
                self.assertNotIn(folder, response.text)
                self.db.refresh(self.job)
                self.assertEqual(self.job.status, JobStatus.FAILED)
                if filename == "company.csv":
                    Path(folder, filename).write_text("transaction_id,amount,status\na,1,success\n")
            response = self.client.get(f"/uploads/{records[1].id}/columns")
            self.assertEqual(response.status_code, 410)
            self.assertNotIn(folder, response.text)

    @classmethod
    def setUpClass(cls):
        cls.password_hash = hash_password("valid-test-password")

    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine, expire_on_commit=False)
        self.db = self.Session()
        rate_db = patch("app.core.throttling.SessionLocal", self.Session)
        rate_db.start()
        self.addCleanup(rate_db.stop)
        self.company = Company(company_name="One", company_email="one@example.com")
        self.other = Company(company_name="Two", company_email="two@example.com")
        self.db.add_all([self.company, self.other])
        self.db.flush()
        self.user = User(company_id=self.company.id, full_name="Test", email="test@example.com",
            password_hash=self.password_hash, role=UserRole.ADMIN, is_active=True)
        self.db.add(self.user)
        self.db.flush()
        self.job = ReconciliationJob(company_id=self.company.id, created_by=self.user.id, job_name="Own", status=JobStatus.COMPLETED)
        self.foreign = ReconciliationJob(company_id=self.other.id, created_by=self.user.id, job_name="Foreign", status=JobStatus.COMPLETED)
        self.db.add_all([self.job, self.foreign])
        self.db.commit()
        def database():
            with self.Session() as db:
                yield db
        app.dependency_overrides[get_db] = database
        app.dependency_overrides[require_active_subscription] = active_user
        self.client = TestClient(app, raise_server_exceptions=False)
        self.sign_in_cookie()

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        self.db.close()
        self.engine.dispose()

    def sign_in_cookie(self):
        token = create_access_token({"sub": str(self.user.id), "ver": self.user.session_version})
        self.client.cookies.set(SESSION_COOKIE, token)

    def csrf(self):
        response = self.client.get("/auth/csrf")
        self.assertEqual(response.status_code, 200)
        return {"X-CSRF-Token": response.json()["csrf_token"]}

    def add_results(self):
        for number in range(7):
            self.db.add(ReconciliationResult(company_id=self.company.id, job_id=self.job.id,
                transaction_id=f"txn-{number}", status=ReconciliationStatus.MATCHED if number % 2 == 0 else ReconciliationStatus.MISSING_IN_COMPANY))
        self.db.add(ReconciliationResult(company_id=self.other.id, job_id=self.foreign.id,
            transaction_id="secret", status=ReconciliationStatus.MATCHED))
        self.db.commit()

    def test_upload_list_matches_mapping_without_redirect_and_is_tenant_scoped(self):
        from app.models.upload import Upload, UploadType, UploadStatus
        own = Upload(company_id=self.company.id, job_id=self.job.id, uploaded_by=self.user.id,
            original_filename="company.csv", stored_filename="company.csv", storage_path="local://company.csv",
            file_size=20, checksum="a" * 64, mime_type="text/csv",
            upload_type=UploadType.COMPANY, status=UploadStatus.UPLOADED)
        foreign = Upload(company_id=self.other.id, job_id=self.foreign.id, uploaded_by=self.user.id,
            original_filename="private.csv", stored_filename="private.csv", storage_path="local://private.csv",
            file_size=20, checksum="b" * 64, mime_type="text/csv",
            upload_type=UploadType.COMPANY, status=UploadStatus.UPLOADED)
        self.db.add_all([own, foreign])
        self.db.commit()
        mapped = self.client.get(f"/uploads/job/{self.job.id}")
        self.assertEqual(mapped.status_code, 200)
        for path in ("/uploads", "/uploads/"):
            response = self.client.get(path, follow_redirects=False)
            self.assertEqual(response.status_code, 200)
            self.assertNotIn("location", response.headers)
            self.assertEqual(response.json(), mapped.json())
            self.assertEqual([row["id"] for row in response.json()], [str(own.id)])
            self.assertNotIn("storage_path", response.json()[0])

    def test_cross_company_run_denied_without_mutation_or_publish(self):
        with patch("app.workers.reconciliation_worker.run_reconciliation_task.apply_async") as publish:
            with self.assertRaises(HTTPException) as failure:
                ReconciliationService.enqueue(self.db, self.company.id, self.foreign.id, self.user.id)
            self.assertEqual(failure.exception.status_code, 404)
            self.assertEqual(self.foreign.status, JobStatus.COMPLETED)
            publish.assert_not_called()

    def test_worker_failure_does_not_change_foreign_job(self):
        from app.workers.reconciliation_worker import run_reconciliation_task
        with patch("app.workers.reconciliation_worker.SessionLocal", self.Session):
            run_reconciliation_task.run(str(self.company.id), str(self.foreign.id), str(self.user.id), "wrong")
        self.db.refresh(self.foreign)
        self.assertEqual(self.foreign.status, JobStatus.COMPLETED)

    def test_duplicate_enqueue_and_publish_failure(self):
        with patch.object(settings, "RECONCILIATION_MODE", "async"):
            with patch("app.workers.reconciliation_worker.run_reconciliation_task.apply_async"):
                ReconciliationService.enqueue(self.db, self.company.id, self.job.id, self.user.id)
                with self.assertRaises(HTTPException) as failure:
                    ReconciliationService.enqueue(self.db, self.company.id, self.job.id, self.user.id)
                self.assertEqual(failure.exception.status_code, 409)
            self.job.status = JobStatus.FAILED
            self.db.commit()
            with patch("app.core.redis_health.circuit.available", return_value=True), patch("app.workers.reconciliation_worker.run_reconciliation_task.apply_async", side_effect=RuntimeError("offline")):
                result = ReconciliationService.enqueue(self.db, self.company.id, self.job.id, self.user.id)
            self.db.refresh(self.job)
            self.assertEqual(result["result_state"], "QUEUED")
            self.assertEqual(self.job.status, JobStatus.QUEUED)

    def test_pagination_filter_summary_and_tenant_isolation(self):
        self.add_results()
        first = self.client.get(f"/reconciliation/{self.job.id}/results?page_size=2").json()
        second = self.client.get(f"/reconciliation/{self.job.id}/results?page_size=2&page=2").json()
        self.assertEqual(first["total"], 7)
        self.assertEqual(first["matched"], 4)
        self.assertEqual(len(first["results"]), 2)
        self.assertTrue({r["id"] for r in first["results"]}.isdisjoint(r["id"] for r in second["results"]))
        filtered = self.client.get(f"/reconciliation/{self.job.id}/results?status=MISSING&search=txn-1").json()
        self.assertEqual(filtered["filtered_total"], 1)
        self.assertEqual(filtered["total"], 7)
        self.assertEqual(self.client.get(f"/reconciliation/{self.foreign.id}/results").status_code, 404)
        self.assertEqual(self.client.get(f"/reconciliation/{self.job.id}/results?page_size=201").status_code, 422)
        self.assertEqual(self.client.get(f"/reconciliation/{self.job.id}/results?status=INVALID").status_code, 422)
        self.assertEqual(self.client.get(f"/reconciliation/{self.job.id}/results?search=%25").json()["filtered_total"], 0)

    def test_disabled_account_and_reenabled_session_revocation(self):
        UserService.set_active(self.db, self.user, self.user.id, False)
        self.assertEqual(self.client.get(f"/reconciliation/{self.job.id}/results").status_code, 401)
        with self.assertRaises(HTTPException):
            AuthService.login(self.db, SimpleNamespace(email=self.user.email, password="valid-test-password"))
        UserService.set_active(self.db, self.user, self.user.id, True)
        self.assertEqual(self.client.get(f"/reconciliation/{self.job.id}/results").status_code, 401)

    def test_reset_revokes_sessions_and_token_is_single_use(self):
        reset = PasswordResetToken(user_id=self.user.id, token_hash=hashlib.sha256(b"reset").hexdigest(), expires_at=datetime.utcnow()+timedelta(minutes=5))
        self.db.add(reset)
        self.db.commit()
        AuthService.reset_password(self.db, "reset", "replacement-password")
        self.assertEqual(self.client.get(f"/reconciliation/{self.job.id}/results").status_code, 401)
        with self.assertRaises(ValueError):
            AuthService.reset_password(self.db, "reset", "replacement-password")

    def test_csrf_and_logout_revocation(self):
        self.assertEqual(self.client.post("/auth/logout").status_code, 403)
        old = self.client.cookies.get(SESSION_COOKIE)
        self.assertEqual(self.client.post("/auth/logout", headers=self.csrf()).status_code, 200)
        self.client.cookies.set(SESSION_COOKIE, old)
        self.assertEqual(self.client.get(f"/reconciliation/{self.job.id}/results").status_code, 401)

    def test_login_only_sets_httponly_cookie(self):
        self.client.cookies.clear()
        response = self.client.post("/auth/login", headers=self.csrf(), json={"email": self.user.email, "password": "valid-test-password"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotIn("access_token", response.json())
        self.assertIn("HttpOnly", response.headers["set-cookie"])
        self.assertIn("SameSite=lax", response.headers["set-cookie"])

    def test_password_policy_consistent(self):
        from app.schemas.company import CompanyRegistration
        from app.schemas.user import CreateUserRequest
        from app.schemas.auth import PasswordResetConfirm
        from pydantic import ValidationError
        for model, payload in [(CompanyRegistration, dict(company_name="A", company_email="a@example.com", phone="1", address="x", admin_name="A", admin_email="a@example.com")),
                (CreateUserRequest, dict(full_name="A", email="a@example.com", role="Finance")),
                (PasswordResetConfirm, dict(token="x"))]:
            with self.assertRaises(ValidationError):
                model(password="short", **payload)
            model(password="valid-long-password", **payload)

    def test_export_generation_expiry_and_ownership(self):
        self.add_results()
        with patch("app.workers.export_worker.build_export.delay"):
            response = self.client.post(f"/reconciliation/{self.job.id}/export/excel?status=MATCHED", headers=self.csrf())
        self.assertEqual(response.status_code, 202, response.text)
        from uuid import UUID
        export_id = UUID(response.json()["id"])
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)):
            generate_export(self.db, export_id, self.company.id)
            record = self.db.get(ExportRequest, export_id)
            self.assertEqual(record.state, "READY")
            self.assertEqual(record.row_count, 4)
            self.assertEqual(self.client.get(f"/reconciliation/exports/{export_id}/download").status_code, 200)
            record.company_id = self.other.id
            self.db.commit()
            self.assertEqual(self.client.get(f"/reconciliation/exports/{export_id}").status_code, 404)
            record.expires_at = datetime.utcnow()-timedelta(seconds=1)
            self.db.commit()
            reference = record.storage_path
            expire_exports(self.db)
            self.assertFalse(storage.local_path(reference).exists())

    def test_csv_and_excel_exports_are_ready_for_download(self):
        self.add_results()
        expected = {
            "csv": ("text/csv", ".csv", b"transaction_id,company_amount"),
            "excel": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx", b"PK"),
        }
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)):
            for export_format, (content_type, extension, signature) in expected.items():
                with self.subTest(format=export_format), patch("app.workers.export_worker.build_export.delay"):
                    response = self.client.post(
                        f"/reconciliation/{self.job.id}/export/{export_format}?status=MATCHED",
                        headers=self.csrf(),
                    )
                self.assertEqual(response.status_code, 202, response.text)
                export_id = UUID(response.json()["id"])
                generate_export(self.db, export_id, self.company.id)

                status = self.client.get(f"/reconciliation/exports/{export_id}")
                self.assertEqual(status.status_code, 200, status.text)
                self.assertEqual(status.json()["state"], "READY")
                self.assertEqual(status.json()["row_count"], 4)

                download = self.client.get(f"/reconciliation/exports/{export_id}/download")
                self.assertEqual(download.status_code, 200, download.text)
                self.assertTrue(download.headers["content-type"].startswith(content_type))
                self.assertIn(extension, download.headers["content-disposition"])
                self.assertIn(signature, download.content)

    def test_unexpected_error_is_sanitized(self):
        with patch("app.services.dashboard_service.DashboardService.company_summary", side_effect=RuntimeError("private database password")):
            response = self.client.get("/dashboard/summary")
        self.assertEqual(response.status_code, 500)
        self.assertNotIn("private database password", response.text)
        self.assertIn("application/problem+json", response.headers["content-type"])

    def test_user_responses_exclude_credentials(self):
        response = self.client.get("/users/")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotIn("password_hash", response.text)
        self.assertNotIn("session_version", response.text)


class FileBoundaryTests(unittest.TestCase):
    def test_s3_missing_key_is_distinct_from_permission_failure(self):
        from botocore.exceptions import ClientError
        with patch.object(settings, "S3_BUCKET", "test-bucket"), patch.object(settings, "S3_PREFIX", "recon"), patch.object(storage, "s3_client") as client:
            client.return_value.get_object.side_effect = ClientError({"Error": {"Code": "NoSuchKey"}}, "GetObject")
            with self.assertRaises(FileNotFoundError):
                with storage.materialize("s3://test-bucket/recon/missing.csv"):
                    self.fail("Missing object should not produce a file")
            client.return_value.get_object.side_effect = ClientError({"Error": {"Code": "AccessDenied"}}, "GetObject")
            with self.assertRaises(ClientError):
                with storage.materialize("s3://test-bucket/recon/private.csv"):
                    self.fail("Permission failure should not produce a file")

    def test_railway_local_storage_requires_volume_and_sync(self):
        from app.core.config import Settings, BASE_DIR
        options = dict(_env_file=None, ENVIRONMENT="production", STORAGE_BACKEND="local",
            RECONCILIATION_MODE="sync", REDIS_URL=None,
            SECRET_KEY="test-only-secret-with-at-least-32-bytes", FRONTEND_URL="https://example.com")
        with patch.dict(os.environ, {"RAILWAY_ENVIRONMENT_ID": "test", "RAILWAY_VOLUME_MOUNT_PATH": str(BASE_DIR / "storage")}):
            self.assertFalse(Settings(**options).database_echo)
            with self.assertRaisesRegex(RuntimeError, "volume"):
                Settings(**{**options, "RECONCILIATION_MODE": "async"})
            for mount in ("", str(BASE_DIR / "unrelated"), str(BASE_DIR / "storage" / "uploads")):
                with patch.dict(os.environ, {"RAILWAY_VOLUME_MOUNT_PATH": mount}):
                    with self.assertRaisesRegex(RuntimeError, "volume"):
                        Settings(**options)

    def test_csv_content_and_dimension_limits(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"sample.csv"
            path.write_text("id,amount\na,1\nb,2\n")
            validate_file(path)
            with patch.object(settings, "MAX_SPREADSHEET_ROWS", 1), self.assertRaises(ValueError):
                validate_file(path)
            path.write_bytes(b"PK\x00binary")
            with self.assertRaises(ValueError):
                validate_file(path)

    def test_xlsx_expansion_limit(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"sample.xlsx"
            book = Workbook()
            book.active.append(["id", "amount"])
            book.active.append(["a", 1])
            book.save(path)
            book.close()
            validate_file(path)
            with patch.object(settings, "MAX_XLSX_EXPANDED_BYTES", 10), self.assertRaises(ValueError):
                validate_file(path)

    def test_local_storage_roundtrip_and_traversal(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)/"store"), patch.object(settings, "STORAGE_BACKEND", "local"):
            source = Path(folder)/"sample.csv"
            source.write_text("id\na\n")
            reference = storage.put_file(source, "uploads/company/sample.csv")
            with storage.materialize(reference) as local:
                self.assertEqual(local.read_bytes(), source.read_bytes())
            with self.assertRaises(ValueError):
                storage.local_path("local://../outside")
            storage.delete_file(reference)

    def test_formula_neutralization_and_export_bounds(self):
        from app.services.export_writers import COLUMNS
        row = {key: None for key in COLUMNS}
        row.update(transaction_id="=HYPERLINK(\"bad\")", company_amount=-2, status="MATCHED")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"export.xlsx"
            write_excel(iter([row]), path)
            book = load_workbook(path, read_only=True)
            cell = list(book.active.rows)[1][0]
            self.assertEqual(cell.data_type, "s")
            self.assertTrue(cell.value.startswith("'="))
            book.close()
            self.assertEqual(safe_cell(-2), -2)
            pdf = Path(folder)/"export.pdf"
            self.assertEqual(write_pdf(iter([row]*90), pdf), 90)
            self.assertIn(b"/Count 3", pdf.read_bytes())
            with patch.object(settings, "EXPORT_MAX_BYTES", 10), self.assertRaises(ValueError):
                write_excel(iter([row]), path)

    def test_csrf_bound_to_session_and_secure_cookie(self):
        response = Response()
        token = issue_csrf(response, "session-a")
        request = Request({"type": "http", "headers": [(b"x-csrf-token", token.encode()),
            (b"cookie", f"{CSRF_COOKIE}={token}; {SESSION_COOKIE}=session-b".encode())]})
        with self.assertRaises(HTTPException):
            validate_csrf(request)
        with patch.object(settings, "ENVIRONMENT", "production"):
            response = Response()
            set_session(response, "session")
            self.assertIn("Secure", response.headers["set-cookie"])

    def test_durable_rate_limit_survives_redis_outage(self):
        from app.core.throttling import throttle
        engine = create_engine("sqlite://")
        Base.metadata.create_all(engine)
        self.addCleanup(engine.dispose)
        with patch("app.core.throttling.SessionLocal", sessionmaker(bind=engine)), patch.object(settings, "REDIS_URL", "redis://unavailable"):
            throttle("test", "one", 2)
            throttle("test", "one", 2)
            with self.assertRaises(HTTPException) as failure:
                throttle("test", "one", 2)
            self.assertEqual(failure.exception.status_code, 429)
        from sqlalchemy.exc import OperationalError
        with patch("app.core.throttling.SessionLocal", side_effect=OperationalError("offline", {}, None)):
            with self.assertRaises(HTTPException) as failure:
                throttle("test", "two", 2)
            self.assertEqual(failure.exception.status_code, 503)
