"""Bound bytes, spreadsheet expansion and dimensions before processing."""
import csv
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from defusedxml.ElementTree import iterparse
from app.core.config import settings


def validate_file(path):
    path = Path(path)
    if not path.stat().st_size or path.stat().st_size > settings.MAX_FILE_BYTES:
        raise ValueError("File is empty or exceeds the upload size limit.")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            with path.open(encoding="utf-8-sig", newline="") as source:
                for number, row in enumerate(csv.reader(source, strict=True)):
                    if number > settings.MAX_SPREADSHEET_ROWS:
                        raise ValueError("File exceeds the row limit.")
                    if len(row) > settings.MAX_SPREADSHEET_COLUMNS or any("\0" in cell for cell in row):
                        raise ValueError("Invalid CSV content or too many columns.")
        except (UnicodeError, csv.Error) as exc:
            raise ValueError("Upload a valid UTF-8 CSV file.") from exc
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > settings.MAX_XLSX_EXPANDED_BYTES:
                    raise ValueError("Spreadsheet expanded size exceeds the limit.")
                names = {e.filename for e in entries}
                if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                    raise ValueError("File is not an XLSX workbook.")
                for entry in entries:
                    if entry.flag_bits & 1 or (entry.compress_size and entry.file_size / entry.compress_size > 200):
                        raise ValueError("Encrypted or excessively compressed spreadsheets are unsupported.")
                    if "vbaproject" in entry.filename.lower() or "externallinks/" in entry.filename.lower():
                        raise ValueError("Macros and external workbook links are unsupported.")
                    if entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml"):
                        rows = cells = 0
                        with archive.open(entry) as sheet:
                            for _, element in iterparse(sheet, events=("end",)):
                                tag = element.tag.rsplit("}", 1)[-1]
                                if tag == "c":
                                    cells += 1
                                    if cells > settings.MAX_SPREADSHEET_COLUMNS:
                                        raise ValueError("Spreadsheet exceeds the column limit.")
                                elif tag == "row":
                                    rows += 1
                                    cells = 0
                                    if rows > settings.MAX_SPREADSHEET_ROWS + 1:
                                        raise ValueError("Spreadsheet exceeds the row limit.")
                                element.clear()
            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book:
                    if (sheet.max_row or 0) > settings.MAX_SPREADSHEET_ROWS + 1 or (sheet.max_column or 0) > settings.MAX_SPREADSHEET_COLUMNS:
                        raise ValueError("Spreadsheet dimensions exceed the limit.")
            finally:
                book.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Upload a valid, unencrypted XLSX workbook.") from exc
    elif suffix == ".xls":
        import xlrd
        try:
            book = xlrd.open_workbook(path, on_demand=True)
            try:
                for sheet in book.sheets():
                    if sheet.nrows > settings.MAX_SPREADSHEET_ROWS + 1 or sheet.ncols > settings.MAX_SPREADSHEET_COLUMNS:
                        raise ValueError("Spreadsheet dimensions exceed the limit.")
            finally:
                book.release_resources()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Upload a valid, unencrypted XLS workbook.") from exc
    else:
        raise ValueError("Unsupported file type.")
